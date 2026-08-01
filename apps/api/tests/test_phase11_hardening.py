"""
Phase 11.1 — Business Workflow Hardening tests.

HTTP-level integration tests for batch/dashboard/export routes (previously
smoke-tested only via import checks in test_phase11_batch.py), plus the
email connector's disabled-by-default guarantee.
"""
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from unittest.mock import AsyncMock, patch

import openpyxl
import pytest
from httpx import ASGITransport, AsyncClient
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from invoice_pipeline.api.main import app
from invoice_pipeline.db.models import Base, Batch, Document, Invoice, LineItem, Vendor, Workspace
from invoice_pipeline.db.session import get_session
from invoice_pipeline.llm.base import ExtractionMeta
from invoice_pipeline.schemas import FieldValue
from invoice_pipeline.schemas import Invoice as ExtractedInvoice

FIXTURES = Path(__file__).parent / "fixtures"


# ── Shared fixtures ─────────────────────────────────────────────────────────────


@pytest.fixture
async def db_engine():
    engine = create_async_engine("sqlite+aiosqlite:///:memory:", echo=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    yield engine
    await engine.dispose()


@pytest.fixture
async def db_session(db_engine):
    factory = async_sessionmaker(db_engine, expire_on_commit=False)
    async with factory() as session:
        yield session


@pytest.fixture
async def ws(db_session: AsyncSession) -> Workspace:
    """A real guest workspace — every scoped route requires X-Workspace-Id now."""
    w = Workspace(workspace_type="guest", expires_at=datetime.now(timezone.utc) + timedelta(hours=1))
    db_session.add(w)
    await db_session.commit()
    return w


@pytest.fixture
async def client(db_session: AsyncSession, ws: Workspace):
    async def override_get_session():
        yield db_session

    app.dependency_overrides[get_session] = override_get_session
    async with AsyncClient(
        transport=ASGITransport(app=app),
        base_url="http://test",
        headers={"X-Workspace-Id": ws.id},
    ) as ac:
        yield ac
    app.dependency_overrides.clear()


@pytest.fixture
def mock_llm():
    invoice = ExtractedInvoice(
        invoice_number=FieldValue(value="INV-2024-001", confidence=0.95, evidence="INV-2024-001"),
        invoice_date=FieldValue(value="2024-01-15", confidence=0.9, evidence="2024-01-15"),
        vendor_name=FieldValue(value="Acme Corp", confidence=0.92, evidence="Acme Corp"),
        subtotal=FieldValue(value="1000.00", confidence=0.93, evidence="$1000.00"),
        tax_amount=FieldValue(value="100.00", confidence=0.91, evidence="$100.00"),
        total_amount=FieldValue(value="1100.00", confidence=0.96, evidence="$1100.00"),
        currency=FieldValue(value="USD", confidence=0.99, evidence="USD"),
    )
    meta = ExtractionMeta(
        provider_name="openai", model_name="gpt-4o-mini", latency_ms=1.0,
        tokens_in=10, tokens_out=10, cost_estimate=0.0,
    )
    mock_provider = AsyncMock()
    mock_provider.extract = AsyncMock(return_value=(invoice, meta))
    with patch("invoice_pipeline.llm.factory._provider_instance", mock_provider):
        with patch("invoice_pipeline.stages.field_extract.get_provider", return_value=mock_provider):
            yield mock_provider


def upload_dir(workspace_id: str) -> Path:
    from invoice_pipeline.utils.storage import upload_dir as _upload_dir

    return _upload_dir(workspace_id)


@pytest.fixture(autouse=True)
def patch_batch_background_session(db_engine):
    """batch.py's background task opens its own session (the request-scoped one
    is closed by the time it runs) via the module-level async_session_factory.
    Point that at this test's in-memory engine instead of the real configured DB."""
    from sqlalchemy.ext.asyncio import async_sessionmaker

    test_factory = async_sessionmaker(db_engine, expire_on_commit=False)
    with patch("invoice_pipeline.api.routes.batch.async_session_factory", test_factory):
        yield


# ── Batch upload: non-terminating + accurate stats ──────────────────────────────


@pytest.mark.asyncio
async def test_batch_upload_mixed_success_and_skip_is_non_terminating(client, mock_llm, ws):
    good_bytes = (FIXTURES / "sample_invoice.pdf").read_bytes()
    files = [
        ("files", ("good.pdf", good_bytes, "application/pdf")),
        ("files", ("blocked.exe", b"MZ\x90\x00", "application/x-msdownload")),
    ]
    res = await client.post(
        "/batch/upload", files=files, headers={"X-Workspace-Id": ws.id}
    )
    assert res.status_code == 202
    body = res.json()
    assert body["total_files"] == 2
    assert body["status"] == "processing"

    # Background task runs to completion within the same ASGI call in tests.
    detail = (
        await client.get(f"/batch/{body['batch_id']}", headers={"X-Workspace-Id": ws.id})
    ).json()
    assert detail["total_files"] == 2
    assert detail["skipped"] == 1
    assert detail["completed"] + detail["failed"] == 1
    # Skipped files never become Document rows; only processed ones do.
    assert len(detail["documents"]) == detail["completed"] + detail["failed"]


@pytest.mark.asyncio
async def test_batch_upload_corrupted_pdf_does_not_abort_other_files(client, mock_llm, ws):
    good_bytes = (FIXTURES / "sample_invoice.pdf").read_bytes()
    corrupted = b"%PDF-1.4 not actually a valid pdf structure" + bytes(range(200))
    files = [
        ("files", ("first_good.pdf", good_bytes, "application/pdf")),
        ("files", ("corrupted.pdf", corrupted, "application/pdf")),
    ]
    res = await client.post(
        "/batch/upload", files=files, headers={"X-Workspace-Id": ws.id}
    )
    assert res.status_code == 202
    body = res.json()
    assert body["total_files"] == 2

    detail = (
        await client.get(f"/batch/{body['batch_id']}", headers={"X-Workspace-Id": ws.id})
    ).json()
    filenames = {d["filename"] for d in detail["documents"]}
    # Both files got processed — the corrupted one never aborted the batch.
    assert filenames == {"first_good.pdf", "corrupted.pdf"}
    # Every result lands in exactly one bucket — none silently dropped.
    assert detail["completed"] + detail["failed"] + detail["skipped"] == 2
    
    # Assert errors round-trip for the failed document
    corrupted_doc = next(d for d in detail["documents"] if d["filename"] == "corrupted.pdf")
    assert corrupted_doc["status"] == "failed"
    assert "errors" in corrupted_doc
    assert isinstance(corrupted_doc["errors"], list)
    assert len(corrupted_doc["errors"]) > 0


@pytest.mark.asyncio
async def test_batch_upload_moderate_scale_does_not_crash(client, mock_llm, ws):
    """Stand-in for a 1000+ file folder: exercises the same per-file loop with
    no N-dependent branching, at a scale that stays fast in CI."""
    files = [("files", (f"invoice_{i}.exe", b"MZ", "application/x-msdownload")) for i in range(40)]
    res = await client.post(
        "/batch/upload", files=files, headers={"X-Workspace-Id": ws.id}
    )
    assert res.status_code == 202
    body = res.json()
    assert body["total_files"] == 40

    detail = (
        await client.get(f"/batch/{body['batch_id']}", headers={"X-Workspace-Id": ws.id})
    ).json()
    assert detail["skipped"] == 40


@pytest.mark.asyncio
async def test_batch_upload_empty_file_list_is_rejected(client, ws):
    # `files` is a required list[UploadFile] — FastAPI 422s before the handler's
    # own "if not files" guard is ever reached, making that guard dead code.
    res = await client.post(
        "/batch/upload", files=[], headers={"X-Workspace-Id": ws.id}
    )
    assert res.status_code == 422


# ── Batch retry: only failed documents ───────────────────────────────────────────


@pytest.mark.asyncio
async def test_retry_failed_only_retries_failed_documents(client, db_session, mock_llm, ws):
    batch = Batch(
        workspace_id=ws.id, upload_source="test", total_files=2, completed=1, failed=1,
        pending=0, skipped=0,
    )
    db_session.add(batch)
    await db_session.flush()

    good_bytes = (FIXTURES / "sample_invoice.pdf").read_bytes()
    ok_doc = Document(
        id="a" * 64, workspace_id=ws.id, filename="ok.pdf", mime_type="application/pdf",
        file_size_bytes=len(good_bytes), status="complete", errors=[], batch_id=batch.id,
    )
    failed_doc = Document(
        id="b" * 64, workspace_id=ws.id, filename="retry_me.pdf", mime_type="application/pdf",
        file_size_bytes=len(good_bytes), status="failed", errors=[{"stage": "pipeline", "message": "boom"}],
        batch_id=batch.id,
    )
    db_session.add_all([ok_doc, failed_doc])
    await db_session.commit()

    updir = upload_dir(ws.id)
    (updir / failed_doc.id).write_bytes(good_bytes)
    try:
        res = await client.post(
            f"/batch/{batch.id}/retry-failed", headers={"X-Workspace-Id": ws.id}
        )
        assert res.status_code == 202
        body = res.json()
        assert body["retried"] == 1  # only the failed doc, not the complete one
        assert body["succeeded"] == 1
    finally:
        (updir / failed_doc.id).unlink(missing_ok=True)


@pytest.mark.asyncio
async def test_list_batches_pagination_uses_real_limit_offset(client, db_session, ws):
    for i in range(3):
        db_session.add(Batch(id=f"batch-{i}", workspace_id=ws.id, upload_source="test", total_files=1))
    await db_session.commit()

    res = await client.get(
        "/batch/", params={"skip": 0, "limit": 2}, headers={"X-Workspace-Id": ws.id}
    )
    body = res.json()
    assert body["total"] == 3  # total reflects all rows, not just the page
    assert len(body["batches"]) == 2

    res = await client.get(
        "/batch/", params={"skip": 2, "limit": 2}, headers={"X-Workspace-Id": ws.id}
    )
    body = res.json()
    assert len(body["batches"]) == 1


@pytest.mark.asyncio
async def test_retry_failed_no_failed_docs_is_a_noop(client, db_session, ws):
    batch = Batch(
        workspace_id=ws.id, upload_source="test", total_files=1, completed=1, failed=0,
        pending=0, skipped=0,
    )
    db_session.add(batch)
    await db_session.commit()

    res = await client.post(
        f"/batch/{batch.id}/retry-failed", headers={"X-Workspace-Id": ws.id}
    )
    assert res.status_code == 202
    assert res.json()["retried"] == 0


# ── Dashboard search: filters actually apply ─────────────────────────────────────


async def _seed_invoice(
    session: AsyncSession, workspace_id: str, *, doc_id: str, invoice_number: str, vendor_name: str,
    confidence: float, needs_review: bool,
) -> None:
    vendor = Vendor(id=f"v-{doc_id}", workspace_id=workspace_id, canonical_name=vendor_name, aliases=[])
    session.add(vendor)
    doc = Document(
        id=doc_id, workspace_id=workspace_id, filename=f"{invoice_number}.pdf",
        mime_type="application/pdf",
        file_size_bytes=10, status="needs_review" if needs_review else "complete", errors=[],
    )
    session.add(doc)
    await session.flush()
    inv = Invoice(
        workspace_id=workspace_id, document_id=doc_id, vendor_id=vendor.id, invoice_number=invoice_number,
        invoice_date="2024-01-15", total_amount=Decimal("500.00"), currency="USD",
        needs_review=needs_review, raw_extraction={},
        confidence_breakdown={"overall_score": confidence},
    )
    session.add(inv)
    await session.commit()


@pytest.mark.asyncio
async def test_dashboard_search_filters_by_vendor_and_min_confidence(client, db_session, ws):
    await _seed_invoice(
        db_session, ws.id, doc_id="d" * 64, invoice_number="INV-A", vendor_name="Acme Corp",
        confidence=0.95, needs_review=False,
    )
    await _seed_invoice(
        db_session, ws.id, doc_id="e" * 64, invoice_number="INV-B", vendor_name="Beta LLC",
        confidence=0.40, needs_review=True,
    )

    res = await client.get("/dashboard/search", params={"vendor": "Acme"})
    body = res.json()
    assert body["total"] == 1
    assert body["items"][0]["invoice_number"] == "INV-A"

    res = await client.get("/dashboard/search", params={"min_confidence": 0.9})
    body = res.json()
    assert body["total"] == 1
    assert body["items"][0]["vendor_name"] == "Acme Corp"

    res = await client.get("/dashboard/search", params={"q": "INV-B"})
    body = res.json()
    assert body["total"] == 1
    assert body["items"][0]["invoice_number"] == "INV-B"

    res = await client.get("/dashboard/search", params={"status": "needs_review"})
    body = res.json()
    assert body["total"] == 1
    assert body["items"][0]["invoice_number"] == "INV-B"


@pytest.mark.asyncio
async def test_dashboard_stats_counts_match_seeded_rows(client, db_session, ws):
    await _seed_invoice(
        db_session, ws.id, doc_id="f" * 64, invoice_number="INV-C", vendor_name="Gamma Inc",
        confidence=0.8, needs_review=False,
    )
    res = await client.get("/dashboard/stats")
    body = res.json()
    assert body["totals"]["invoices"] == 1
    assert body["totals"]["approved"] == 1
    assert len(body["vendor_statistics"]) == 1
    assert body["vendor_statistics"][0]["vendor_name"] == "Gamma Inc"


# ── Excel export: content matches DB, Decimal precision preserved ────────────────


@pytest.mark.asyncio
async def test_export_excel_totals_match_db_and_includes_line_items(client, db_session, ws):
    doc = Document(
        id="c" * 64, workspace_id=ws.id, filename="export_me.pdf", mime_type="application/pdf",
        file_size_bytes=10, status="complete", errors=[],
    )
    db_session.add(doc)
    await db_session.flush()
    inv = Invoice(
        workspace_id=ws.id, document_id=doc.id, invoice_number="INV-EXPORT", invoice_date="2024-03-01",
        subtotal=Decimal("999.9999"), tax_amount=Decimal("10.0001"), total_amount=Decimal("1010.0000"),
        currency="EUR", needs_review=False, raw_extraction={},
    )
    db_session.add(inv)
    await db_session.flush()
    db_session.add(
        LineItem(
            invoice_id=inv.id, position=0, description_raw="Widget",
            quantity_raw="2", unit_price_raw="5.00", total_raw="10.00",
        )
    )
    await db_session.commit()

    res = await client.get("/export/excel")
    assert res.status_code == 200
    assert res.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    import io

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws1 = wb["Invoices"]
    header = [c.value for c in ws1[1]]
    row = {header[i]: ws1[2][i].value for i in range(len(header))}
    assert row["Invoice Number"] == "INV-EXPORT"
    # Written value must equal the DB Decimal exactly — no float rounding drift.
    assert Decimal(str(row["Subtotal"])) == Decimal("999.9999")
    assert Decimal(str(row["Total"])) == Decimal("1010.0000")

    ws2 = wb["Line Items"]
    assert ws2[2][0].value == "INV-EXPORT"
    assert ws2[2][2].value == "Widget"


@pytest.mark.asyncio
async def test_export_history_recorded(client, db_session):
    await client.get("/export/excel")
    res = await client.get("/export/history")
    assert res.status_code == 200
    assert len(res.json()["exports"]) == 1


# ── Email connector: disabled by default, zero effect ────────────────────────────


def test_email_import_disabled_by_default():
    from invoice_pipeline.config import settings

    assert settings.EMAIL_IMPORT_ENABLED is False


def test_email_connector_refuses_to_connect_when_disabled():
    from invoice_pipeline.email_connector.connector import EmailConnector

    conn = EmailConnector(host="imap.example.com", port=993, username="u", password="p")
    with patch("imaplib.IMAP4_SSL") as mock_imap:
        with pytest.raises(RuntimeError, match="disabled"):
            conn.connect()
        # No network call was ever attempted — truly zero effect when disabled.
        mock_imap.assert_not_called()
